import json
import logging
from typing import List, Dict,Optional
from langgraph.prebuilt import ToolNode
from langchain_core.messages import SystemMessage, AnyMessage, HumanMessage
from dotenv import load_dotenv
from pydantic import BaseModel, Field
from colorama import init, Fore, Back, Style
from fastapi import WebSocket,WebSocketDisconnect,HTTPException
import asyncio

# docx related imports
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

'''LOCAL IMPORTS'''
from schemas import State
from prompts.prompts import *

'''IMPORT ALL TOOLS HERE AND CREATE LIST OF TOOLS TO BE PASSED TO THE AGENT.'''
from tools.query_chromadb import query_chromadb
from pathlib import Path
import importlib.util
import certifi
import os

# Dynamically resolve the path to config.py
CURRENT_FILE = Path(__file__).resolve()
SAGAN_MULTIMODAL = CURRENT_FILE.parent.parent.parent.parent
CONFIG_PATH = SAGAN_MULTIMODAL / "config.py"

os.environ['SSL_CERT_FILE'] = certifi.where()

# Load config.py dynamically
spec = importlib.util.spec_from_file_location("config", CONFIG_PATH)
config = importlib.util.module_from_spec(spec)
spec.loader.exec_module(config)

load_dotenv()
init()

research_tools = [query_chromadb]

'''LLM TO USE'''
from smolagents import ToolCallingAgent, HfApiModel, CodeAgent
# select model
# model_id = "Qwen/Qwen2.5-Coder-32B-Instruct"
model_id = "deepseek-ai/DeepSeek-R1-Distill-Qwen-32B"
# model_id = "meta-llama/Llama-3.3-70B-Instruct"
# model_id = "Qwen/Qwen2.5-72B-Instruct"
# model_id = "mistralai/Mistral-7B-Instruct-v0.3"
# model_id = "NousResearch/Hermes-3-Llama-3.1-8B"
model = HfApiModel(model_id=model_id)  

message_queues = {}

class WebSocketManager:
    def __init__(self):
        self.active_connections = {}
        self.lock = asyncio.Lock()

    async def connect(self, websocket: WebSocket, session_id: str):
        """Accept and track a WebSocket connection."""
        # async with self.lock:
        if session_id in self.active_connections:
            
            # await self.disconnect(session_id)
            print("connect to seesion id already exist")
            await websocket.accept()
            self.active_connections[session_id] = websocket
            self.lock = asyncio.Lock()

            message_queues[session_id] = asyncio.Queue()
            return
        await websocket.accept()
        self.active_connections[session_id] = websocket
        message_queues[session_id] = asyncio.Queue()
        print(f"Connected: {session_id}")

    async def disconnect(self, session_id: str):
        """Remove a disconnected WebSocket."""
        async with self.lock:
            if session_id in self.active_connections:
                self.active_connections.pop(session_id)
                del message_queues[session_id]
                print(f"Disconnected: {session_id}")

    # async def send_message(self, session_id: str, message: str):
    #     """Send a message to a specific client."""
    #     async with self.lock:
    #         websocket = self.active_connections.get(session_id)
    #         if websocket:
    #             try:
    #                 await websocket.send_text(message)
    #                 print(f"Sent to {session_id}: {message}")
    #             except Exception as e:
    #                 print(f"Error sending to {session_id}: {e}")
    #         else:
    #             print(f"Session {session_id} not connected.")
    async def send_message(self, session_id: str, message: str):
        """Send a message to a specific client."""
        # async with self.lock:
        websocket = self.active_connections.get(session_id)
        if websocket:
            try:
                json_message = json.dumps(message)
                await websocket.send_text(json_message)
                print(f"Sent to {session_id}: {json_message}")
            except Exception as e:
                print(f"Error sending to {session_id}: {e}")
        else:
            print(f"Session {session_id} not connected.")

    async def receive_message(self, session_id: str):
        """Receive a message from a specific client."""
        async with self.lock:
            websocket = self.active_connections.get(session_id)
            print(websocket,"websocket 117")
            if websocket:
                try:
                    message = await websocket.receive_text()
                    message_data = json.loads(message)
                    key = message_data.get("key")
                    value = message_data.get("value")
                    print("Received key:", key, "Received value:", value)
                    await message_queues[session_id].put((key, value))
                    print(f"Received from {session_id}: {message_data}")
                    return message_data
                except WebSocketDisconnect:
                    print(f"Error receiving from {session_id}: WebSocket disconnected")
                    await self.disconnect(session_id)  # Clean up
                    return None
                except Exception as e:
                    print(f"Error receiving from {session_id}: {e}")
                    return None
            else:
                print(f"Session {session_id} not connected.")
                await ws_manager.active_connections.clear()
                await websocket.active_connections.clear()
                return None
            
    async def get_message(self, session_id: str, key: str, timeout: int = 500):
        queue = message_queues.get(session_id)
        print(queue,"queue 141")
        if not queue:
            raise HTTPException(status_code=404, detail="Session not found.")
        
        try:
            # Wait for a message with a timeout
            while True:
                key_value = await asyncio.wait_for(queue.get(), timeout=timeout)
                print(key_value,"key_value 150")
                if key_value[0] == key:
                    return key_value[1]
                # Put back unmatched keys
                await queue.put(key_value)
        except asyncio.TimeoutError:
            raise HTTPException(status_code=408, detail=f"No message received within {timeout} seconds.")
ws_manager = WebSocketManager()

async def research_query_generator(state: State) -> State:
    """
    Node to generate research queries and allow user modification via WebSocket.
    """
    session_id = "1234"
    print(f"{Fore.YELLOW}################ RESEARCH QUERY GENERATOR BEGIN #################")
    empty_prompt = """ """
    
    agent = ToolCallingAgent(
        model=model, 
        tools=[], 
        prompt_templates={"system_prompt": empty_prompt}
    )
    
    user_prompt = f"""
    Section Title: {state['section_title']}
    Section Text: {state['section_text']}
    User Prompt: {state['user_prompt']}
    """

    combined_prompt = RESEARCH_QUERY_GENERATOR_PROMPT + "\n" + user_prompt

    try:
        # Get response from agent
        response = agent.provide_final_answer(combined_prompt, images=None)
        
        # Debug logging
        print("\nRaw response from agent:")
        print("-" * 50)
        print(response)
        print("-" * 50)
        
        # Try to clean the response if it contains extra text
        try:
            if '{' in response:
                json_start = response.find('{')
                json_end = response.rfind('}') + 1
                response = response[json_start:json_end]
            
            response_content = json.loads(response)
        except json.JSONDecodeError as e:
            print(f"JSON parsing error: {e}")
            print("Falling back to empty research queries")
            response_content = {"research_queries": []}
        
        # Extract research queries with fallback
        research_queries = response_content.get('research_queries', [])
        if not isinstance(research_queries, list):
            print("Warning: research_queries is not a list, converting to empty list")
            research_queries = []
            
        print("Generated Research Queries:", research_queries)

        # Save the queries to state before attempting WebSocket communication
        state["research_queries"] = research_queries
        state["research_needed"] = bool(research_queries)

        try:
            # WebSocket communication attempt
            session_id = "1234"

            if research_queries:
                await ws_manager.send_message("1234", {
                    "type": "queries",
                    "data": research_queries
                })

            # await ws_manager.send_message("1234", {
            #     "type": "queries",
            #     "data": research_queries
            # })
            
            if bool(research_queries) is False:
                return state


            await ws_manager.send_message("1234", {
                "type": "question1",
                "data": "Would you like to modify or add queries? (yes/no)"
            })
            user_input_str = await ws_manager.get_message(session_id, 'question1')
            
            if user_input_str and user_input_str.lower() == 'yes':
                modified_queries = []
                modified_queries = research_queries
                
                # Modify existing queries
                for i, query in enumerate(research_queries, start=1):
                    # print('do this later')
                    await ws_manager.send_message("1234",{
                        "type":"question",
                        "data":f"Modify query {i} (or leave blank to keep it unchanged):"
                    })
                    new_query = await ws_manager.wait_for_response(
                        "1234"
                    )
                    modified_queries.append(new_query or query)
                
                # Add new queries
                while True:
                    
                    await ws_manager.send_message("1234",{
                        "type":"question2",
                        "data":"Would you like to add a new query? (yes/no)"
                    })
                    add_more= await ws_manager.get_message(session_id, 'question2')
                    # add_more = await ws_manager.wait_for_response(
                    #     "1234"
                    # )

                    print("351",add_more)
                    # add_more = await ws_manager.query_user(
                    #     session_id, "Would you like to add a new query? (yes/no)"
                    # )
                    if add_more and add_more.lower() != 'yes':
                        break
                    await ws_manager.send_message(session_id,{
                        "type":"question3",
                        "data":f"Enter new query {len(modified_queries) + 1}:"
                    })

                    new_query = await ws_manager.get_message(session_id,"question3")
                    print(new_query,"new_query")
                    # new_query = await ws_manager.query_user(
                    #     session_id, f"Enter new query {len(modified_queries) + 1}:"
                    # )

                    if new_query:
                        modified_queries.append(new_query)

                research_queries = modified_queries
            # ------------------------------------------------

            # Update state
            state["research_needed"] = bool(research_queries)
            state["research_queries"] = research_queries
            print(f"Final Research Queries: {research_queries}")

            return state

        except Exception as ws_error:
            print(f"WebSocket communication error: {ws_error}")
            # Continue with existing queries if WebSocket fails
            pass

        return state

    except Exception as e:
        print(f"Error in research_query_generator: {str(e)}")
        state["messages"] = [str(e)]
        state["research_needed"] = False
        state["research_queries"] = []
        return state

# new code below be aware llm code ahead

# def research_query_generator(state: dict, research_queries: List[str]) -> dict:
#     """
#     Process the state and research queries synchronously.
#     This function assumes that all necessary inputs have been provided beforehand.
#     """
#     try:
#         # Update the state with the finalized research queries
#         state["research_needed"] = len(research_queries) > 0
#         state["research_queries"] = research_queries

#         # Log the updated state
#         print("\nFinal research queries:")
#         for i, query in enumerate(research_queries, start=1):
#             print(f"{i}. {query}")

#         return state

#     except Exception as e:
#         print(f"Error in research_query_generator: {str(e)}")
#         state["messages"] = [str(e)]
#         state["research_needed"] = False
#         state["research_queries"] = []
#         return state

# my code ahead
def research_query_answerer(state: State) -> State:
    """
    Takes the generated queries and executes them against the vector database.
    Only runs if research_needed is True.
    """
    print(f"{Fore.BLUE}################ RESEARCH QUERY ANSWERER BEGIN #################")

    if not state.get("research_needed"):
        state["context"] = []
        return state

    context = []
    for query in state["research_queries"]:
        result = query_chromadb(
            str(config.VECTOR_DB_PATHS['astro_db']),  # Use path from config
            config.MODEL_SETTINGS['SENTENCE_TRANSFORMER'],  # Use model setting from config
            query
        )
        context.extend(result)

    state["context"] = context

    print(f"################ QUERY ANSWERER END #################{Style.RESET_ALL}")
    return state


def formatter(state: State):
    print(f"{Fore.LIGHTGREEN_EX}################ FORMATTING NODE BEGIN #################")
    empty_prompt = """ """
    agent = ToolCallingAgent(
        model=model, 
        tools=[], 
        prompt_templates={"system_prompt": empty_prompt}
    )

    # Ensure state values are not None and convert to string if needed
    section_title = str(state.get("section_title", ""))
    user_prompt = str(state.get("user_prompt", ""))
    section_text = str(state.get("section_text", ""))
    context = "\n".join(state.get("context", [])) if isinstance(state.get("context"), list) else str(state.get("context", ""))

    # Modify the formatter prompt to be more specific about maintaining structure
    formatter_user_prompt = f"""
    You are tasked with modifying a section of text while preserving its overall structure.

    Section Title: {section_title}  
    
    Original Section Text:
    {section_text}
    
    User's Modification Request:
    {user_prompt}
    
    Additional Context:
    {context}
    
    Instructions:
    1. Keep all existing subsections intact
    2. Only modify the specific parts relevant to the user's request
    3. Maintain the same formatting and structure as the original
    4. Return the COMPLETE section with your modifications
    5. Your response should be in JSON format with the following structure:
    {{
        "modified_section_text": "entire section text with modifications",
        "ai_message": "brief description of changes made"
    }}
    """

    combined_prompt = FORMATTER_PROMPT + "\n" + formatter_user_prompt
    raw_response = agent.provide_final_answer(combined_prompt, images=None)

    # Debug logging
    print("\nRaw response from agent:")
    print("-" * 50)
    print(raw_response)
    print("-" * 50)

    try:
        # First try to extract just the JSON part if present
        if '{' in raw_response:
            json_start = raw_response.find('{')
            json_end = raw_response.rfind('}') + 1
            json_str = raw_response[json_start:json_end]
            response_data = json.loads(json_str)
        else:
            # If no JSON found, try to parse the whole response
            response_data = json.loads(raw_response)
            
    except json.JSONDecodeError as e:
        print(f"Error parsing JSON response: {e}")
        print("Raw response:", raw_response)
        # Create a more informative fallback response
        response_data = {
            'modified_section_text': state.get('section_text', ''),
            'ai_message': f'Error: Could not process the text modification. Raw response: {raw_response[:100]}...'
        }
    
    # Extract the fields with fallback values
    modified_section_text = response_data.get('modified_section_text', section_text)
    ai_message = response_data.get('ai_message', 'No message provided')

    # Update state with the extracted values
    state['modified_section_text'] = modified_section_text
    state['ai_message'] = ai_message

    print(f"################ FORMATTING NODE END #################{Style.RESET_ALL}")
    return state


# original code below 
# def human_input_node(state: State):
#     print(f"{Fore.LIGHTMAGENTA_EX}################ HUMAN INPUT NODE BEGIN #################")

#     response = input("Saves Changes to the section text? (yes/no): ")
#     state["user_approval"] = response
#     print(f"################ HUMAN INPUT NODE END #################{Style.RESET_ALL}")
#     return state

# gpt code belwo 

async def human_input_node(state: State):
    print(f"{Fore.LIGHTMAGENTA_EX}################ HUMAN INPUT NODE BEGIN #################")
    await ws_manager.send_message("1234",{
        "type":"question4",
        "data":"Saves Changes to the section text? (yes/no): "
    })

    response = await ws_manager.get_message("1234","question4")
    print(response,"response 655")
    # response = input("Saves Changes to the section text? (yes/no): ")

    state["user_approval"] = response
    print(f"################ HUMAN INPUT NODE END #################{Style.RESET_ALL}")
    return state

# helper functions for docx editing.
def insert_paragraph_after(paragraph, text=""):
    """
    Insert a new paragraph immediately after the given paragraph and set its text.
    """
    # Create a new XML element for a paragraph
    new_p_element = OxmlElement("w:p")
    # Insert the new element right after the current paragraph's element
    paragraph._p.addnext(new_p_element)
    # Wrap the new element in a Paragraph object
    new_paragraph = Paragraph(new_p_element, paragraph._parent)
    # Add a run with the provided text
    new_paragraph.add_run(text)
    return new_paragraph

def fill_section_with_text(doc_path: str, section_title: str, section_text: str):
    """
    Open the document at 'doc_path', find the paragraph containing section_title,
    and replace all content after it until the next section with section_text.
    
    Args:
        doc_path (str): Path to the Word document
        section_title (str): Title of the section to modify
        section_text (str): New text to replace the section content with
    """
    # Load the document
    doc = Document(doc_path)
    
    # Find the section and its content
    section_start = None
    for i, paragraph in enumerate(doc.paragraphs):
        if section_title.lower() in paragraph.text.lower():
            section_start = i
            break
    
    if section_start is None:
        print(f"Section '{section_title}' not found in document")
        return
    
    # Find the end of this section (next heading or document end)
    section_end = len(doc.paragraphs)
    for i in range(section_start + 1, len(doc.paragraphs)):
        if doc.paragraphs[i].style.name.startswith('Heading'):
            section_end = i
            break
    
    # Remove all paragraphs between section start and end (except the heading)
    for i in range(section_end - 1, section_start, -1):
        p = doc.paragraphs[i]._element
        p.getparent().remove(p)
    
    # Insert the new text after the section heading
    insert_paragraph_after(doc.paragraphs[section_start], section_text)
    
    # Save the modified document back to the same file
    doc.save(doc_path)


def save_changes(state: State):
    print(f"{Fore.CYAN}################ SAVING CHANGES NODE BEGIN #################")
    
    # Only save if user approved
    state["user_approval"] = "yes" # line added adhoc to test the workflow
    if state.get("user_approval") == 'yes':
        docx_file_path = state.get("rough_draft_path")

        # Get required values from state
        section_title = state.get("section_title")
        modified_text = state.get("modified_section_text")
            
        # Use helper function to update the document
        fill_section_with_text(
            doc_path=docx_file_path,
            section_title=section_title,
            section_text=modified_text
        )
        
        print("Changes saved successfully to DOCX file!")
    else:
        print("User disapproved changes. No changes saved.")

    print(f"################ SAVING CHANGES NODE END #################{Style.RESET_ALL}")
    return state


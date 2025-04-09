from langchain_core.messages import SystemMessage, HumanMessage, AIMessage, ToolMessage
from typing import Dict, Any
import datetime
from pathlib import Path
import json
import os
import sys
import shutil
import importlib.util
from types import ModuleType

# python-docx related imports
from docx import Document
from docx.shared import Pt
from docxtpl import DocxTemplate
import os
import shutil
from copy import deepcopy

# Local imports
from schemas import State

# helper function to load the config.py file dynamically
def load_config_file() -> ModuleType:
    # Dynamically resolve the path to config.py
    CURRENT_FILE = Path(__file__).resolve()
    project_root = CURRENT_FILE.parent.parent.parent.parent
    CONFIG_PATH = project_root / "config.py"
    # Load config.py dynamically
    spec = importlib.util.spec_from_file_location("config", CONFIG_PATH)
    config = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(config)
    return config

# Get the path to the spaider_agent_temp directory
CURRENT_DIR = Path(os.path.dirname(os.path.abspath(__file__)))
SPAIDER_AGENT_TEMP_DIR = CURRENT_DIR.parent
sys.path.append(str(SPAIDER_AGENT_TEMP_DIR))

configfile = load_config_file()

# getting project_id from cookie.json
CURRENT_FILE = Path(__file__).resolve()
project_root = CURRENT_FILE.parent.parent.parent.parent

# getting project_id from cookie.json
with open(project_root / "cookie.json", "r") as f:
    cookie_data = json.load(f)
    current_project_id = cookie_data.get("project_id")


def print_state(state):
    """Print the state in a human-readable format."""
    print("STATE:")
    print("=" * 40)
    for key, value in state.items():
        if key == "messages":
            print(f"{key.upper()}:")
            for msg in value:
                print(f"  - {(msg.type).upper()}: {msg.content}")
        else:
            print(f"{key.upper()}: {value}")
    print("=" * 40)


def serialize_message(msg):
    """Helper function to serialize a message object."""
    if isinstance(msg, (SystemMessage, HumanMessage, AIMessage, ToolMessage)):
        serialized = {
            'type': msg.type,
            'content': msg.content,
            'additional_kwargs': msg.additional_kwargs,
        }
        # Add response_metadata for AIMessage
        if isinstance(msg, AIMessage) and hasattr(msg, 'response_metadata'):
            serialized['response_metadata'] = msg.response_metadata
        return serialized
    return str(msg)

# helper fn for save_state_for_testing
def serialize_state(state: State, node_name: str) -> Dict[str, Any]:
    """
    Serialize the state object into a JSON-compatible dictionary.
    Handles special cases like message objects and includes metadata.
    """
    serialized = {
        'metadata': {
            'node_name': node_name,
            'timestamp': str(datetime.datetime.now()),
        },
        'state': {}
    }
    
    for key, value in state.items():
        if key == 'messages':
            serialized['state'][key] = [serialize_message(msg) for msg in value]
        elif isinstance(value, (list, dict, str, int, float, bool, type(None))):
            serialized['state'][key] = value
        else:
            # For other types, convert to string representation
            serialized['state'][key] = str(value)
    
    return serialized


def save_state_for_testing(state: State, node_name: str):
    """
    Save the state in .json format so it can be reused.
    """
    configfile = load_config_file()

    current_project_id = configfile.get_current_project_id()
    updated_paths = configfile.update_project_paths(current_project_id)
    
    # Use the path from config
    output_dir = updated_paths['nodewise_output']
    
    output_dir.mkdir(parents=True, exist_ok=True)
    
    # Save machine-readable JSON version
    json_path = output_dir / f"{node_name}_state.json"
    serialized_state = serialize_state(state, node_name)
    with json_path.open('w', encoding='utf-8') as f:
        json.dump(serialized_state, f, indent=2, ensure_ascii=False)


# reconstructs the state for a given node, given its .json file in outputpdf/nodewise_output
def load_state_for_testing(node_name: str, output_dir: Path = None) -> State:
    """
    Load a previously saved state for testing purposes.
    Reconstructs the full State object with proper message types and metadata.
    """
    
    json_path = output_dir / f"{node_name}_state.json"
    if not json_path.exists():
        raise FileNotFoundError(f"No saved state found for node: {node_name}")
    
    with json_path.open('r', encoding='utf-8') as f:
        data = json.load(f)
    
    state_dict = data['state']
    
    # Reconstruct message objects with full metadata
    if 'messages' in state_dict:
        messages = []
        for msg_data in state_dict['messages']:
            msg_type = msg_data['type']
            content = msg_data['content']
            additional_kwargs = msg_data.get('additional_kwargs', {})
            response_metadata = msg_data.get('response_metadata', {})
            
            # Create appropriate message object based on type
            if msg_type == 'system':
                msg = SystemMessage(
                    content=content,
                    additional_kwargs=additional_kwargs
                )
            elif msg_type == 'human':
                msg = HumanMessage(
                    content=content,
                    additional_kwargs=additional_kwargs
                )
            elif msg_type == 'ai':
                # For AI messages, include response_metadata
                msg = AIMessage(
                    content=content,
                    additional_kwargs=additional_kwargs,
                    response_metadata=response_metadata
                )
            elif msg_type == 'tool':
                msg = ToolMessage(
                    content=content,
                    additional_kwargs=additional_kwargs
                )
            else:
                continue  # Skip unknown message types
                
            messages.append(msg)
        
        state_dict['messages'] = messages
    
    # Create and return State object with all fields from the original state
    return State(**state_dict)

# Helper function to copy figures
def copy_figures(source_dir: Path, dest_dir: Path):
    """Copy all figures from the source directory to the destination directory."""
    for item in source_dir.glob("figures/*"):
        if item.is_file():
            shutil.copy(item, dest_dir / item.name)


def create_wp_table_in_doc(doc, wp_number, title, num_tasks=3, num_deliverables=3, num_milestones=2, num_researchers=3):
    """
    Create a work package table in the given document object.
    
    Parameters:
    - doc: Document object to add the table to
    - wp_number: Work package number
    - title: Work package title
    - num_tasks: Number of tasks placeholders
    - num_deliverables: Number of deliverable placeholders
    - num_milestones: Number of milestone placeholders
    - num_researchers: Number of researcher rows
    """
    # Add title
    title_para = doc.add_paragraph()
    title_run = title_para.add_run(f"Work Package {wp_number}: {title}")
    title_run.font.size = Pt(14)
    title_run.bold = True
    
    # Create main table with 4 columns (we'll need 4 for the dates row)
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    
    # Basic info rows - WP number
    row = table.rows[0]
    cell = row.cells[0]
    cell.text = "WP number"
    cell.paragraphs[0].runs[0].bold = True
    
    cell = row.cells[1]
    cell.merge(row.cells[2])
    cell.merge(row.cells[3])
    cell.text = str(wp_number)
    cell.paragraphs[0].runs[0].bold = True
    
    # WP title
    row = table.add_row()
    cell = row.cells[0]
    cell.text = "WP title"
    cell.paragraphs[0].runs[0].bold = True
    
    cell = row.cells[1]
    cell.merge(row.cells[2])
    cell.merge(row.cells[3])
    cell.text = title
    cell.paragraphs[0].runs[0].bold = True
    
    # WP leader
    row = table.add_row()
    cell = row.cells[0]
    cell.text = "WP leader"
    cell.paragraphs[0].runs[0].bold = True
    
    cell = row.cells[1]
    cell.merge(row.cells[2])
    cell.merge(row.cells[3])
    cell.text = f"{{{{ wp_leader_{wp_number} }}}}"
    cell.paragraphs[0].runs[0].bold = True
    
    # Date row 
    row = table.add_row()
    cell = row.cells[0]
    cell.text = "Start date"
    cell.paragraphs[0].runs[0].bold = True
    
    cell = row.cells[1]
    cell.text = f"{{{{ start_date_{wp_number} }}}}"
    
    cell = row.cells[2]
    cell.text = "End date"
    cell.paragraphs[0].runs[0].bold = True
    
    cell = row.cells[3]
    cell.text = f"{{{{ end_date_{wp_number} }}}}"
    
    # Objective row
    row = table.add_row()
    cell = row.cells[0]
    cell.text = "Objective"
    cell.paragraphs[0].runs[0].bold = True
    cell.merge(row.cells[1])
    cell.merge(row.cells[2])
    cell.merge(row.cells[3])
    
    # Objective content
    row = table.add_row()
    cell = row.cells[0]
    cell.merge(row.cells[1])
    cell.merge(row.cells[2])
    cell.merge(row.cells[3])
    cell.text = f"{{{{ objective_{wp_number} }}}}"
    
    # Tasks header
    row = table.add_row()
    cell = row.cells[0]
    cell.text = "Tasks"
    cell.paragraphs[0].runs[0].bold = True
    cell.merge(row.cells[1])
    cell.merge(row.cells[2])
    cell.merge(row.cells[3])
    
    # Task items
    for i in range(num_tasks):
        row = table.add_row()
        cell = row.cells[0]
        cell.merge(row.cells[1])
        cell.merge(row.cells[2])
        cell.merge(row.cells[3])
        para = cell.paragraphs[0]
        run = para.add_run(f"• Task {wp_number}.{i+1}. ")
        run.bold = True
        para.add_run(f"{{{{ tasks_{wp_number}[{i}].title }}}}. {{{{ tasks_{wp_number}[{i}].description }}}}")
    
    # Interdependence header
    row = table.add_row()
    cell = row.cells[0]
    cell.text = "Interdependence with other work packages"
    cell.paragraphs[0].runs[0].bold = True
    cell.merge(row.cells[1])
    cell.merge(row.cells[2])
    cell.merge(row.cells[3])
    
    # Interdependence content
    row = table.add_row()
    cell = row.cells[0]
    cell.merge(row.cells[1])
    cell.merge(row.cells[2])
    cell.merge(row.cells[3])
    cell.text = f"{{{{ interdependence_{wp_number} }}}}"
    
    # Deliverables and milestones header
    row = table.add_row()
    cell = row.cells[0]
    cell.text = "Deliverables and milestones"
    cell.paragraphs[0].runs[0].bold = True
    cell.merge(row.cells[1])
    cell.merge(row.cells[2])
    cell.merge(row.cells[3])
    
    # Deliverables header
    row = table.add_row()
    cell = row.cells[0]
    cell.text = "Deliverables"
    cell.paragraphs[0].runs[0].bold = True
    cell.merge(row.cells[1])
    cell.merge(row.cells[2])
    cell.merge(row.cells[3])
    
    # Deliverable items
    for i in range(num_deliverables):
        row = table.add_row()
        cell = row.cells[0]
        cell.merge(row.cells[1])
        cell.merge(row.cells[2])
        cell.merge(row.cells[3])
        para = cell.paragraphs[0]
        run = para.add_run(f"• D{wp_number}.{i+1}. ")
        run.bold = True
        para.add_run(f"{{{{ deliverables_{wp_number}[{i}].title }}}} (M{{{{ deliverables_{wp_number}[{i}].month }}}})")
    
    # Milestones header
    row = table.add_row()
    cell = row.cells[0]
    cell.text = "Milestones"
    cell.paragraphs[0].runs[0].bold = True
    cell.merge(row.cells[1])
    cell.merge(row.cells[2])
    cell.merge(row.cells[3])
    
    # Milestone items
    for i in range(num_milestones):
        row = table.add_row()
        cell = row.cells[0]
        cell.merge(row.cells[1])
        cell.merge(row.cells[2])
        cell.merge(row.cells[3])
        para = cell.paragraphs[0]
        run = para.add_run(f"• M{wp_number}.{i+1}. ")
        run.bold = True
        para.add_run(f"{{{{ milestones_{wp_number}[{i}].title }}}} (M{{{{ milestones_{wp_number}[{i}].month }}}})")
    
    # Human resources header
    row = table.add_row()
    cell = row.cells[0]
    cell.text = "Human resources"
    cell.paragraphs[0].runs[0].bold = True
    cell.merge(row.cells[1])
    cell.merge(row.cells[2])
    cell.merge(row.cells[3])
    
    # HR table
    hr_table = doc.add_table(rows=num_researchers + 1, cols=4)
    hr_table.style = 'Table Grid'
    
    # HR table headers
    headers = ["Name of researcher", "Partner", "Qualification level", "Person*months"]
    for i, header in enumerate(headers):
        cell = hr_table.cell(0, i)
        para = cell.paragraphs[0]
        run = para.add_run(header)
        run.bold = True
    
    # HR table rows
    for i in range(1, num_researchers + 1):
        hr_table.cell(i, 0).text = f"{{{{ researchers_{wp_number}[{i-1}].name }}}}"
        hr_table.cell(i, 1).text = f"{{{{ researchers_{wp_number}[{i-1}].partner }}}}"
        hr_table.cell(i, 2).text = f"{{{{ researchers_{wp_number}[{i-1}].qualification }}}}"
        hr_table.cell(i, 3).text = f"{{{{ researchers_{wp_number}[{i-1}].months }}}}"

def append_wp_table(existing_doc_path, wp_number, title, output_path, **kwargs):
    """
    Append a work package table to an existing document.
    
    Parameters:
    - existing_doc_path: Path to existing document
    - wp_number: Work package number
    - title: Work package title
    - output_path: Path to save the output document
    - kwargs: Additional parameters (num_tasks, num_deliverables, etc.)
    """
    doc = Document(existing_doc_path)
    doc.add_page_break()
    create_wp_table_in_doc(doc, wp_number, title, **kwargs)
    doc.save(output_path)

def replace_wp_markers_preserving_format(input_doc_path, output_doc_path, wp_info_list):
    """
    Replace WP markers while preserving document formatting, headers, and footers.
    """
    # Make a copy of the original document
    shutil.copy2(input_doc_path, output_doc_path)
    
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.shared import Pt
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    # Border XML for directly setting borders
    border_xml = f'''<w:tblBorders {nsdecls("w")}>
        <w:top w:val="single" w:sz="4" w:space="0" w:color="auto"/>
        <w:left w:val="single" w:sz="4" w:space="0" w:color="auto"/>
        <w:bottom w:val="single" w:sz="4" w:space="0" w:color="auto"/>
        <w:right w:val="single" w:sz="4" w:space="0" w:color="auto"/>
        <w:insideH w:val="single" w:sz="4" w:space="0" w:color="auto"/>
        <w:insideV w:val="single" w:sz="4" w:space="0" w:color="auto"/>
      </w:tblBorders>'''

    # Open document once and process all markers
    doc = Document(output_doc_path)
    
    # Sort WP list by marker position to avoid index shifting problems
    # First find all marker positions
    marker_positions = {}
    for i, para in enumerate(doc.paragraphs):
        for wp_info in wp_info_list:
            marker = f'<WP{wp_info["number"]}>'
            if marker in para.text:
                marker_positions[wp_info["number"]] = i
    
    # Sort WP list by position in document (process earlier markers first)
    sorted_wp_list = sorted(
        [wp for wp in wp_info_list if wp["number"] in marker_positions],
        key=lambda wp: marker_positions[wp["number"]]
    )
    
    # Keep track of inserted elements to adjust position calculations
    inserted_elements = 0
    
    # Process markers in order they appear in document
    for wp_info in sorted_wp_list:
        wp_num = wp_info['number']
        marker = f'<WP{wp_num}>'
        
        for i, para in enumerate(doc.paragraphs):
            if marker in para.text:
                # Create a temporary document for the WP content
                temp_doc = Document()
                
                # Add the title
                title_para = temp_doc.add_paragraph()
                title_run = title_para.add_run(f"Work Package {wp_num}: {wp_info['title']}")
                title_run.font.size = Pt(14)
                title_run.bold = True
                
                # Create main table with borders
                kwargs = {k: v for k, v in wp_info.items() if k not in ['number', 'title']}
                table = temp_doc.add_table(rows=1, cols=4)
                tbl = table._element.xpath('./w:tblPr')[0]
                tbl_borders = parse_xml(border_xml)
                tbl.append(tbl_borders)
                
                # Populate tables
                create_wp_table_content(temp_doc, table, wp_num, wp_info['title'], border_xml, **kwargs)
                
                # Get the paragraph with marker
                marker_para = para._element
                parent = marker_para.getparent()
                idx = parent.index(marker_para)
                
                # Clear or remove the marker paragraph
                para.text = para.text.replace(marker, "").strip()
                if not para.text:
                    parent.remove(marker_para)
                
                # Insert content from temporary document
                for element in reversed(temp_doc._element.body):
                    if idx < len(parent):
                        parent.insert(idx + 1, deepcopy(element))
                    else:
                        parent.append(deepcopy(element))
                
                break
    
    # Save document once with all changes
    doc.save(output_doc_path)

def create_wp_table_content(doc, table, wp_number, title, border_xml, num_tasks=3, num_deliverables=3, num_milestones=2, num_researchers=3):
    """
    Populate an existing table with work package content and ensure borders are applied.
    """
    from docx.shared import Pt
    from docx.oxml import parse_xml
    
    # Basic info rows - WP number
    row = table.rows[0]
    cell = row.cells[0]
    cell.text = "WP number"
    cell.paragraphs[0].runs[0].bold = True
    
    cell = row.cells[1]
    cell.merge(row.cells[2])
    cell.merge(row.cells[3])
    cell.text = str(wp_number)
    cell.paragraphs[0].runs[0].bold = True
    
    # WP title
    row = table.add_row()
    cell = row.cells[0]
    cell.text = "WP title"
    cell.paragraphs[0].runs[0].bold = True
    
    cell = row.cells[1]
    cell.merge(row.cells[2])
    cell.merge(row.cells[3])
    cell.text = title
    cell.paragraphs[0].runs[0].bold = True
    
    # WP leader
    row = table.add_row()
    cell = row.cells[0]
    cell.text = "WP leader"
    cell.paragraphs[0].runs[0].bold = True
    
    cell = row.cells[1]
    cell.merge(row.cells[2])
    cell.merge(row.cells[3])
    cell.text = f"{{{{ wp_leader_{wp_number} }}}}"
    cell.paragraphs[0].runs[0].bold = True
    
    # Date row 
    row = table.add_row()
    cell = row.cells[0]
    cell.text = "Start date"
    cell.paragraphs[0].runs[0].bold = True
    
    cell = row.cells[1]
    cell.text = f"{{{{ start_date_{wp_number} }}}}"
    
    cell = row.cells[2]
    cell.text = "End date"
    cell.paragraphs[0].runs[0].bold = True
    
    cell = row.cells[3]
    cell.text = f"{{{{ end_date_{wp_number} }}}}"
    
    # Objective row
    row = table.add_row()
    cell = row.cells[0]
    cell.text = "Objective"
    cell.paragraphs[0].runs[0].bold = True
    cell.merge(row.cells[1])
    cell.merge(row.cells[2])
    cell.merge(row.cells[3])
    
    # Objective content
    row = table.add_row()
    cell = row.cells[0]
    cell.merge(row.cells[1])
    cell.merge(row.cells[2])
    cell.merge(row.cells[3])
    cell.text = f"{{{{ objective_{wp_number} }}}}"
    
    # Tasks header
    row = table.add_row()
    cell = row.cells[0]
    cell.text = "Tasks"
    cell.paragraphs[0].runs[0].bold = True
    cell.merge(row.cells[1])
    cell.merge(row.cells[2])
    cell.merge(row.cells[3])
    
    # Task items
    for i in range(num_tasks):
        row = table.add_row()
        cell = row.cells[0]
        cell.merge(row.cells[1])
        cell.merge(row.cells[2])
        cell.merge(row.cells[3])
        para = cell.paragraphs[0]
        run = para.add_run(f"• Task {wp_number}.{i+1}. ")
        run.bold = True
        para.add_run(f"{{{{ tasks_{wp_number}[{i}].title }}}}. {{{{ tasks_{wp_number}[{i}].description }}}}")
    
    # Interdependence header
    row = table.add_row()
    cell = row.cells[0]
    cell.text = "Interdependence with other work packages"
    cell.paragraphs[0].runs[0].bold = True
    cell.merge(row.cells[1])
    cell.merge(row.cells[2])
    cell.merge(row.cells[3])
    
    # Interdependence content
    row = table.add_row()
    cell = row.cells[0]
    cell.merge(row.cells[1])
    cell.merge(row.cells[2])
    cell.merge(row.cells[3])
    cell.text = f"{{{{ interdependence_{wp_number} }}}}"
    
    # Deliverables and milestones header
    row = table.add_row()
    cell = row.cells[0]
    cell.text = "Deliverables and milestones"
    cell.paragraphs[0].runs[0].bold = True
    cell.merge(row.cells[1])
    cell.merge(row.cells[2])
    cell.merge(row.cells[3])
    
    # Deliverables header
    row = table.add_row()
    cell = row.cells[0]
    cell.text = "Deliverables"
    cell.paragraphs[0].runs[0].bold = True
    cell.merge(row.cells[1])
    cell.merge(row.cells[2])
    cell.merge(row.cells[3])
    
    # Deliverable items
    for i in range(num_deliverables):
        row = table.add_row()
        cell = row.cells[0]
        cell.merge(row.cells[1])
        cell.merge(row.cells[2])
        cell.merge(row.cells[3])
        para = cell.paragraphs[0]
        run = para.add_run(f"• D{wp_number}.{i+1}. ")
        run.bold = True
        para.add_run(f"{{{{ deliverables_{wp_number}[{i}].title }}}} (M{{{{ deliverables_{wp_number}[{i}].month }}}})")
    
    # Milestones header
    row = table.add_row()
    cell = row.cells[0]
    cell.text = "Milestones"
    cell.paragraphs[0].runs[0].bold = True
    cell.merge(row.cells[1])
    cell.merge(row.cells[2])
    cell.merge(row.cells[3])
    
    # Milestone items
    for i in range(num_milestones):
        row = table.add_row()
        cell = row.cells[0]
        cell.merge(row.cells[1])
        cell.merge(row.cells[2])
        cell.merge(row.cells[3])
        para = cell.paragraphs[0]
        run = para.add_run(f"• M{wp_number}.{i+1}. ")
        run.bold = True
        para.add_run(f"{{{{ milestones_{wp_number}[{i}].title }}}} (M{{{{ milestones_{wp_number}[{i}].month }}}})")
    
    # Human resources header
    row = table.add_row()
    cell = row.cells[0]
    cell.text = "Human resources"
    cell.paragraphs[0].runs[0].bold = True
    cell.merge(row.cells[1])
    cell.merge(row.cells[2])
    cell.merge(row.cells[3])
    
    # HR table
    hr_table = doc.add_table(rows=num_researchers + 1, cols=4)
    
    # Apply borders to HR table
    hr_tbl = hr_table._element.xpath('./w:tblPr')[0]
    hr_borders = parse_xml(border_xml)
    hr_tbl.append(hr_borders)
    
    # HR table headers
    headers = ["Name of researcher", "Partner", "Qualification level", "Person*months"]
    for i, header in enumerate(headers):
        cell = hr_table.cell(0, i)
        para = cell.paragraphs[0]
        run = para.add_run(header)
        run.bold = True
    
    # HR table rows
    for i in range(1, num_researchers + 1):
        hr_table.cell(i, 0).text = f"{{{{ researchers_{wp_number}[{i-1}].name }}}}"
        hr_table.cell(i, 1).text = f"{{{{ researchers_{wp_number}[{i-1}].partner }}}}"
        hr_table.cell(i, 2).text = f"{{{{ researchers_{wp_number}[{i-1}].qualification }}}}"
        hr_table.cell(i, 3).text = f"{{{{ researchers_{wp_number}[{i-1}].months }}}}"

def render_wp_markers_with_context(input_doc_path, output_doc_path, wp_info_list, context_data):
    """
    Replace WP markers and render template variables.
    
    Parameters:
    - input_doc_path: Path to input document
    - output_doc_path: Path to output document
    - wp_info_list: List of dictionaries with WP info
    - context_data: Dictionary with template variables
    """
    temp_file = "temp_with_markers_replaced.docx"
    try:
        # Replace markers with tables preserving formatting
        replace_wp_markers_preserving_format(input_doc_path, temp_file, wp_info_list)
        
        # Render template variables
        template = DocxTemplate(temp_file)
        template.render(context_data)
        template.save(output_doc_path)
    finally:
        if os.path.exists(temp_file):
            os.remove(temp_file)